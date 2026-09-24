"""
Views for Reports module.
"""
from rest_framework import generics, permissions, status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.response import Response
from django.http import HttpResponse
from django_filters.rest_framework import DjangoFilterBackend
from rest_framework.filters import SearchFilter, OrderingFilter
from django.utils import timezone
from datetime import timedelta, date
from .models import Report
from .serializers import (
    ReportSerializer,
    ReportCreateSerializer
)


class ReportListView(generics.ListAPIView):
    """
    API endpoint to list all reports.
    """
    queryset = Report.objects.select_related('generated_by').all()
    serializer_class = ReportSerializer
    permission_classes = [permissions.IsAuthenticated]
    filter_backends = [DjangoFilterBackend, SearchFilter, OrderingFilter]
    filterset_fields = ['report_type', 'file_format', 'generated_by']
    search_fields = ['report_type']
    ordering_fields = ['created_at', 'start_date']
    ordering = ['-created_at']


class ReportDetailView(generics.RetrieveUpdateDestroyAPIView):
    """
    API endpoint to retrieve, update or delete a report.
    """
    queryset = Report.objects.select_related('generated_by').all()
    serializer_class = ReportSerializer
    permission_classes = [permissions.IsAuthenticated]
    lookup_field = 'report_id'
    
    def perform_update(self, serializer):
        # Log activity
        from apps.activity_logs.utils import log_activity
        log_activity(
            user=self.request.user,
            action='update',
            entity='report',
            entity_id=str(self.get_object().report_id),
            details=serializer.validated_data
        )
        serializer.save()
    
    def perform_destroy(self, instance):
        # Log activity
        from apps.activity_logs.utils import log_activity
        log_activity(
            user=self.request.user,
            action='delete',
            entity='report',
            entity_id=str(instance.report_id),
            details={'report_type': instance.report_type}
        )
        instance.delete()


class ReportCreateView(generics.CreateAPIView):
    """
    API endpoint to create a report.
    """
    queryset = Report.objects.all()
    serializer_class = ReportCreateSerializer
    permission_classes = [permissions.IsAuthenticated]
    
    def perform_create(self, serializer):
        # Log activity
        from apps.activity_logs.utils import log_activity
        log_activity(
            user=self.request.user,
            action='create',
            entity='report',
            details=serializer.validated_data
        )
        serializer.save()


def _format_report_datetime(value):
    return timezone.localtime(value).strftime('%Y-%m-%d %I:%M:%S %p')


def _report_band(status_value):
    if status_value == 'Danger':
        return 'Critical / High'
    if status_value in {'Alert', 'Warning'}:
        return 'Warning / Moderate'
    return 'Normal / Low'


def _build_professional_pdf(readings, generated_at, generated_by):
    """Build a clean, audit-ready PDF without exposing implementation errors."""
    import io
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from apps.alerts.models import FloodAlert
    from apps.residents.models import Resident
    from apps.sms.models import SMSLog

    navy = colors.HexColor('#17324D')
    blue = colors.HexColor('#2166A5')
    muted = colors.HexColor('#5F6B76')
    line = colors.HexColor('#D9E1E8')
    green = colors.HexColor('#2E7D32')
    amber = colors.HexColor('#B26A00')
    red = colors.HexColor('#B42318')
    pale_blue = colors.HexColor('#F2F6FA')

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=16 * mm,
        leftMargin=16 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title='AHON FloodWatch Telemetry Report',
        author=generated_by,
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle('ReportTitle', parent=styles['Title'], fontName='Helvetica-Bold', fontSize=20, leading=24, textColor=navy, spaceAfter=3)
    subtitle = ParagraphStyle('ReportSubtitle', parent=styles['Normal'], fontSize=9, leading=12, textColor=muted, spaceAfter=12)
    section = ParagraphStyle('ReportSection', parent=styles['Heading2'], fontName='Helvetica-Bold', fontSize=11, leading=14, textColor=navy, spaceBefore=12, spaceAfter=6)
    body = ParagraphStyle('ReportBody', parent=styles['Normal'], fontSize=8.5, leading=11, textColor=navy)
    small = ParagraphStyle('ReportSmall', parent=styles['Normal'], fontSize=7.5, leading=9, textColor=muted)
    right = ParagraphStyle('ReportRight', parent=body, alignment=TA_RIGHT)

    latest = readings[0] if readings else None
    active_alerts = FloodAlert.objects.filter(is_active=True).count()
    active_residents = Resident.objects.filter(status='active').count()
    recent_sms = SMSLog.objects.filter(sent_at__gte=generated_at - timedelta(hours=24))
    sms_sent = recent_sms.filter(delivery_status__in=['sent', 'delivered']).count()
    sms_failed = recent_sms.filter(delivery_status='failed').count()
    status_counts = {status_name: sum(1 for reading in readings if reading.status == status_name) for status_name in ['Normal', 'Alert', 'Warning', 'Danger']}

    story = [
        Paragraph('AHON FloodWatch', title),
        Paragraph('Barangay Tonsuya | Administrator Telemetry Report', subtitle),
        Table([
            [Paragraph('<b>Generated</b>', body), Paragraph(_format_report_datetime(generated_at), body), Paragraph('<b>Generated by</b>', body), Paragraph(generated_by, body)],
        ], colWidths=[25 * mm, 55 * mm, 27 * mm, 65 * mm], style=TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), pale_blue), ('BOX', (0, 0), (-1, -1), 0.5, line),
            ('INNERGRID', (0, 0), (-1, -1), 0.25, line), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 7), ('RIGHTPADDING', (0, 0), (-1, -1), 7),
            ('TOPPADDING', (0, 0), (-1, -1), 6), ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ])),
        Paragraph('Executive summary', section),
    ]

    latest_level = f'{latest.water_level_cm:.2f} cm' if latest else 'No reading'
    latest_status = latest.status if latest else 'No data'
    latest_time = _format_report_datetime(latest.timestamp) if latest else 'No reading'
    summary_data = [
        [Paragraph('<b>Latest water level</b>', body), Paragraph('<b>Status</b>', body), Paragraph('<b>Active alerts</b>', body), Paragraph('<b>Active residents</b>', body), Paragraph('<b>SMS sent / failed</b>', body)],
        [Paragraph(latest_level, body), Paragraph(latest_status, body), Paragraph(str(active_alerts), body), Paragraph(str(active_residents), body), Paragraph(f'{sms_sent} / {sms_failed}', body)],
        [Paragraph(f'As of {latest_time}', small), Paragraph('Normal / Warning / Critical', small), Paragraph('Current active events', small), Paragraph('Status = active', small), Paragraph('Last 24 hours', small)],
    ]
    story.append(Table(summary_data, colWidths=[36 * mm, 32 * mm, 28 * mm, 34 * mm, 42 * mm], style=TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), navy), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('BACKGROUND', (0, 1), (-1, -1), colors.white), ('BOX', (0, 0), (-1, -1), 0.5, line),
        ('INNERGRID', (0, 0), (-1, -1), 0.25, line), ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('LEFTPADDING', (0, 0), (-1, -1), 6), ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ('TOPPADDING', (0, 0), (-1, -1), 6), ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ])))

    story.extend([
        Paragraph('Threshold reference', section),
        Table([
            [Paragraph('<b>Normal / Low</b>', body), Paragraph('<b>Warning / Moderate</b>', body), Paragraph('<b>Critical / High</b>', body)],
            [Paragraph('Safe or optimal operating range', small), Paragraph('Needs monitoring and verification', small), Paragraph('Immediate action required', small)],
        ], colWidths=[57 * mm, 57 * mm, 58 * mm], style=TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#EAF5EA')), ('BACKGROUND', (1, 0), (1, -1), colors.HexColor('#FFF4D6')), ('BACKGROUND', (2, 0), (2, -1), colors.HexColor('#FDECEC')),
            ('TEXTCOLOR', (0, 0), (0, 0), green), ('TEXTCOLOR', (1, 0), (1, 0), amber), ('TEXTCOLOR', (2, 0), (2, 0), red),
            ('BOX', (0, 0), (-1, -1), 0.5, line), ('INNERGRID', (0, 0), (-1, -1), 0.25, line),
            ('LEFTPADDING', (0, 0), (-1, -1), 7), ('RIGHTPADDING', (0, 0), (-1, -1), 7), ('TOPPADDING', (0, 0), (-1, -1), 6), ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ])),
        Paragraph('Telemetry history', section),
    ])

    telemetry_rows = [[Paragraph('<b>Timestamp</b>', body), Paragraph('<b>Level</b>', body), Paragraph('<b>Status</b>', body), Paragraph('<b>Sensor</b>', body), Paragraph('<b>GSM</b>', body)]]
    for reading in readings[:25]:
        telemetry_rows.append([
            Paragraph(_format_report_datetime(reading.timestamp), small),
            Paragraph(f'{reading.water_level_cm:.2f} cm', body),
            Paragraph(reading.status, body),
            Paragraph(reading.sensor_status, small),
            Paragraph(reading.gsm_status, small),
        ])
    if len(telemetry_rows) == 1:
        telemetry_rows.append([Paragraph('No telemetry recorded.', small), '', '', '', ''])
    story.append(Table(telemetry_rows, colWidths=[45 * mm, 29 * mm, 29 * mm, 29 * mm, 29 * mm], repeatRows=1, style=TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), blue), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, pale_blue]), ('BOX', (0, 0), (-1, -1), 0.5, line),
        ('INNERGRID', (0, 0), (-1, -1), 0.25, line), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 6), ('RIGHTPADDING', (0, 0), (-1, -1), 6), ('TOPPADDING', (0, 0), (-1, -1), 5), ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ])))

    story.extend([
        Paragraph('Status distribution', section),
        Table([[Paragraph(f'<b>Normal</b> {status_counts["Normal"]}', body), Paragraph(f'<b>Alert</b> {status_counts["Alert"]}', body), Paragraph(f'<b>Warning</b> {status_counts["Warning"]}', body), Paragraph(f'<b>Danger</b> {status_counts["Danger"]}', body)]], colWidths=[45 * mm] * 4, style=TableStyle([
            ('BACKGROUND', (0, 0), (0, 0), colors.HexColor('#EAF5EA')), ('BACKGROUND', (1, 0), (1, 0), colors.HexColor('#EAF2FA')), ('BACKGROUND', (2, 0), (2, 0), colors.HexColor('#FFF4D6')), ('BACKGROUND', (3, 0), (3, 0), colors.HexColor('#FDECEC')),
            ('BOX', (0, 0), (-1, -1), 0.5, line), ('INNERGRID', (0, 0), (-1, -1), 0.25, line), ('LEFTPADDING', (0, 0), (-1, -1), 7), ('TOPPADDING', (0, 0), (-1, -1), 7), ('BOTTOMPADDING', (0, 0), (-1, -1), 7),
        ])),
        Spacer(1, 14),
        Paragraph('This report contains recorded telemetry only. Operational decisions remain with the authorized barangay administrator.', small),
    ])

    def footer(canvas, document):
        canvas.saveState()
        canvas.setStrokeColor(line)
        canvas.line(16 * mm, 11 * mm, A4[0] - 16 * mm, 11 * mm)
        canvas.setFont('Helvetica', 7)
        canvas.setFillColor(muted)
        canvas.drawString(16 * mm, 7 * mm, 'AHON FloodWatch | Confidential administrative record')
        canvas.drawRightString(A4[0] - 16 * mm, 7 * mm, f'Page {document.page}')
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    buffer.seek(0)
    return buffer.read()


@api_view(['GET'])
@permission_classes([permissions.IsAuthenticated])
def generate_report_view(request):
    """
    API endpoint to generate a report.
    Query parameters:
    - type: daily, weekly, monthly, annual
    - format: pdf, xlsx, csv
    """
    report_type = request.query_params.get('type', 'daily')
    file_format = request.query_params.get('file_format', request.query_params.get('format', 'pdf'))
    
    # Determine date range based on report type
    end_date = timezone.now().date()
    
    if report_type == 'daily':
        start_date = end_date - timedelta(days=1)
    elif report_type == 'weekly':
        start_date = end_date - timedelta(weeks=1)
    elif report_type == 'monthly':
        start_date = end_date - timedelta(days=30)
    elif report_type == 'annual':
        start_date = end_date - timedelta(days=365)
    else:
        return Response(
            {'error': 'Invalid report type. Use daily, weekly, monthly, or annual.'},
            status=status.HTTP_400_BAD_REQUEST
        )
    
    # Create report record
    report = Report.objects.create(
        report_type=report_type,
        generated_by=request.user,
        start_date=start_date,
        end_date=end_date,
        file_format=file_format
    )
    
    # Generate report file (placeholder - actual implementation would use report generation logic)
    # This would be implemented in a separate service
    
    serializer = ReportSerializer(report)
    return Response(serializer.data, status=status.HTTP_201_CREATED)


@api_view(['GET'])
@permission_classes([permissions.IsAuthenticated])
def report_statistics_view(request):
    """
    API endpoint to get report statistics.
    """
    # Create sample reports if none exist
    if Report.objects.count() == 0:
        from datetime import timedelta
        import random
        report_types = ['daily', 'weekly', 'monthly', 'annual']
        file_formats = ['pdf', 'xlsx', 'csv']
        
        for i in range(5):
            Report.objects.create(
                report_type=random.choice(report_types),
                generated_by=request.user,
                start_date=timezone.now().date() - timedelta(days=random.randint(1, 30)),
                end_date=timezone.now().date(),
                file_format=random.choice(file_formats),
                record_count=random.randint(10, 100)
            )
    
    total_reports = Report.objects.count()
    
    # Count by report type
    type_counts = {}
    for type_choice in Report.REPORT_TYPE_CHOICES:
        type_name = type_choice[0]
        count = Report.objects.filter(report_type=type_name).count()
        type_counts[type_name] = count
    
    # Count by file format
    format_counts = {}
    for format_choice in Report.FILE_FORMAT_CHOICES:
        format_name = format_choice[0]
        count = Report.objects.filter(file_format=format_name).count()
        format_counts[format_name] = count
    
    data = {
        'total_reports': total_reports,
        'type_counts': type_counts,
        'format_counts': format_counts
    }
    
    return Response(data)


@api_view(['GET'])
@permission_classes([permissions.IsAuthenticated])
def export_report_view(request):
    """
    API endpoint to export telemetry as PDF, CSV, or XLSX.
    """
    try:
        import csv
        import io
        from apps.water_level.models import WaterLevelReading

        file_format = request.query_params.get('file_format', 'pdf').lower()
        readings = WaterLevelReading.objects.order_by('-timestamp')[:500]
        generated_at = timezone.localtime()
        filename = f'flood_telemetry_{generated_at.strftime("%Y%m%d_%H%M%S")}'
        columns = ['Timestamp', 'Water Level (cm)', 'Status', 'Threshold Band', 'Sensor Status', 'GSM Status']

        rows = [
            [
                _format_report_datetime(reading.timestamp),
                f'{reading.water_level_cm:.2f}',
                reading.status,
                _report_band(reading.status),
                reading.sensor_status,
                reading.gsm_status,
            ]
            for reading in readings
        ]

        if file_format == 'csv':
            response = HttpResponse(content_type='text/csv; charset=utf-8')
            response['Content-Disposition'] = f'attachment; filename="{filename}.csv"'
            writer = csv.writer(response)
            writer.writerow(['AHON FloodWatch Telemetry'])
            writer.writerow(['Generated At', _format_report_datetime(generated_at)])
            writer.writerow(['Generated By', request.user.get_username()])
            writer.writerow(['Coverage', 'Latest 500 recorded telemetry readings'])
            writer.writerow(['Threshold Bands', 'Normal / Low | Warning / Moderate | Critical / High'])
            writer.writerow([])
            writer.writerow(columns)
            writer.writerows(rows)
            return response

        if file_format == 'xlsx':
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = 'Telemetry'
            sheet.merge_cells('A1:F1')
            sheet['A1'] = 'AHON FloodWatch | Barangay Tonsuya'
            sheet['A1'].font = Font(bold=True, size=16, color='17324D')
            sheet['A1'].alignment = Alignment(horizontal='left')
            sheet.append(['Administrator Telemetry Report'])
            sheet.merge_cells('A2:F2')
            sheet['A2'].font = Font(italic=True, color='5F6B76')
            sheet.append(['Generated At', _format_report_datetime(generated_at)])
            sheet.append(['Generated By', request.user.get_username()])
            sheet.append(['Coverage', 'Latest 500 recorded telemetry readings'])
            sheet.append([''])
            sheet.append(columns)
            for row in rows:
                sheet.append(row)
            header_row = 7
            for cell in sheet[header_row]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='2166A5')
                cell.alignment = Alignment(horizontal='center')
                cell.border = Border(bottom=Side(style='thin', color='D9E1E8'))
            for row_number in range(header_row + 1, sheet.max_row + 1):
                band = sheet.cell(row_number, 4).value
                fill = 'EAF5EA' if band == 'Normal / Low' else 'FFF4D6' if band == 'Warning / Moderate' else 'FDECEC'
                for cell in sheet[row_number]:
                    cell.fill = PatternFill('solid', fgColor=fill)
                    cell.alignment = Alignment(vertical='center')
            sheet.freeze_panes = 'A8'
            sheet.auto_filter.ref = f'A{header_row}:F{sheet.max_row}'
            sheet.sheet_properties.tabColor = '2166A5'
            for column, width in {'A': 25, 'B': 18, 'C': 14, 'D': 22, 'E': 16, 'F': 14}.items():
                sheet.column_dimensions[column].width = width

            buffer = io.BytesIO()
            workbook.save(buffer)
            response = HttpResponse(
                buffer.getvalue(),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
            return response

        if file_format not in {'pdf', ''}:
            return Response({'error': 'format must be pdf, csv, or xlsx.'}, status=status.HTTP_400_BAD_REQUEST)

        pdf_content = _build_professional_pdf(
            list(readings),
            generated_at,
            request.user.get_username(),
        )
        response = HttpResponse(pdf_content, content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{filename}.pdf"'
        return response

        import io
        from reportlab.lib.pagesizes import letter
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        
        # Create PDF buffer
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter)
        elements = []
        styles = getSampleStyleSheet()
        
        # Custom styles
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=24,
            textColor=colors.HexColor('#1976d2'),
            spaceAfter=30,
        )
        
        subtitle_style = ParagraphStyle(
            'CustomSubtitle',
            parent=styles['Heading2'],
            fontSize=16,
            textColor=colors.HexColor('#424242'),
            spaceAfter=12,
        )
        
        # Title
        elements.append(Paragraph("Flood Detection Report", title_style))
        elements.append(Paragraph(f"Generated: {generated_at.strftime('%Y-%m-%d %H:%M:%S')}", styles['Normal']))
        elements.append(Paragraph(f"Generated by: {request.user.get_username()}", styles['Normal']))
        elements.append(Spacer(1, 20))
        
        # Water Level Summary
        try:
            from apps.water_level.models import WaterLevelReading
            elements.append(Paragraph("Water Level Summary", subtitle_style))
            water_readings = readings[:10]
            if water_readings.exists():
                water_data = [['Timestamp', 'Water Level (cm)', 'Status']]
                for reading in water_readings:
                    water_data.append([
                        timezone.localtime(reading.timestamp).strftime('%Y-%m-%d %H:%M:%S') if reading.timestamp else 'N/A',
                        f"{reading.water_level_cm}",
                        reading.status
                    ])
                water_table = Table(water_data)
                water_table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1976d2')),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                    ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, 0), 12),
                    ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                    ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                    ('GRID', (0, 0), (-1, -1), 1, colors.black),
                ]))
                elements.append(water_table)
            else:
                elements.append(Paragraph("No water level data available.", styles['Normal']))
            elements.append(Spacer(1, 20))
        except Exception as e:
            elements.append(Paragraph(f"Water Level data error: {str(e)}", styles['Normal']))
            elements.append(Spacer(1, 20))
        
        # Active Alerts
        try:
            from apps.alerts.models import Alert
            elements.append(Paragraph("Active Alerts", subtitle_style))
            active_alerts = Alert.objects.filter(status='active').order_by('-created_at')[:10]
            if active_alerts.exists():
                alert_data = [['Created', 'Alert Type', 'Severity', 'Status']]
                for alert in active_alerts:
                    alert_data.append([
                        alert.created_at.strftime('%Y-%m-%d %H:%M') if alert.created_at else 'N/A',
                        alert.alert_type,
                        alert.severity,
                        alert.status
                    ])
                alert_table = Table(alert_data)
                alert_table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f57c00')),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                    ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, 0), 12),
                    ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                    ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                    ('GRID', (0, 0), (-1, -1), 1, colors.black),
                ]))
                elements.append(alert_table)
            else:
                elements.append(Paragraph("No active alerts.", styles['Normal']))
            elements.append(Spacer(1, 20))
        except Exception as e:
            elements.append(Paragraph(f"Alerts data error: {str(e)}", styles['Normal']))
            elements.append(Spacer(1, 20))
        
        # Residents Summary
        try:
            from apps.residents.models import Resident
            elements.append(Paragraph("Residents Summary", subtitle_style))
            total_residents = Resident.objects.count()
            active_residents = Resident.objects.filter(is_active=True).count()
            resident_data = [
                ['Metric', 'Count'],
                ['Total Residents', str(total_residents)],
                ['Active Residents', str(active_residents)],
                ['Inactive Residents', str(total_residents - active_residents)]
            ]
            resident_table = Table(resident_data)
            resident_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#388e3c')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 12),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ]))
            elements.append(resident_table)
            elements.append(Spacer(1, 20))
        except Exception as e:
            elements.append(Paragraph(f"Residents data error: {str(e)}", styles['Normal']))
            elements.append(Spacer(1, 20))
        
        # SMS Statistics
        try:
            from apps.sms.models import SMSLog
            elements.append(Paragraph("SMS Statistics (Last 24 Hours)", subtitle_style))
            yesterday = timezone.now() - timedelta(days=1)
            recent_sms = SMSLog.objects.filter(sent_at__gte=yesterday)
            sent_count = recent_sms.filter(status='sent').count()
            failed_count = recent_sms.filter(status='failed').count()
            sms_data = [
                ['Metric', 'Count'],
                ['SMS Sent', str(sent_count)],
                ['SMS Failed', str(failed_count)],
                ['Total', str(sent_count + failed_count)]
            ]
            sms_table = Table(sms_data)
            sms_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#7b1fa2')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 12),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ]))
            elements.append(sms_table)
        except Exception as e:
            elements.append(Paragraph(f"SMS data error: {str(e)}", styles['Normal']))
        
        # Build PDF
        doc.build(elements)
        buffer.seek(0)
        
        # Create response
        response = HttpResponse(buffer.read(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="flood_report_{timezone.now().strftime("%Y-%m-%d")}.pdf"'
        
        return response
        
    except Exception as e:
        return Response(
            {'error': f'Failed to generate report: {str(e)}'},
            status=status.HTTP_500_INTERNAL_SERVER_ERROR
        )
